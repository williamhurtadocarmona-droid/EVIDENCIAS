import streamlit as st
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import os
from datetime import datetime

st.set_page_config(
    page_title="Informe de Evidencias SENA", 
    page_icon="📷", 
    layout="centered"
)

st.title("📷 Generador de Informe de Actividades y Evidencias")
st.write("Diligencia los grupos, carga las fotografías y genera el informe en Word automáticamente.")

# Posibles nombres de la plantilla en el repositorio
OPCIONES_PLANTILLA = [
    "INFORME DE ACTIVIDADES Y EVIDENCIAS.docx",
    "INFORME_DE_ACTIVIDADES_Y_EVIDENCIAS.docx"
]

# Detectar cuál archivo existe realmente en GitHub
PLANTILLA_WORD = None
for opc in OPCIONES_PLANTILLA:
    if os.path.exists(opc):
        PLANTILLA_WORD = opc
        break

MESES = {
    1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril",
    5: "Mayo", 6: "Junio", 7: "Julio", 8: "Agosto",
    9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre"
}

hoy = datetime.now()
fecha_hoy_str = f"{hoy.day} de {MESES[hoy.month]} de {hoy.year}"

st.info(f"📅 **Fecha del Informe:** {fecha_hoy_str}")

if not PLANTILLA_WORD:
    st.error("⚠️ **No se encontró el archivo Word de la plantilla en GitHub.**")
    st.warning("Por favor sube a GitHub el archivo `INFORME DE ACTIVIDADES Y EVIDENCIAS.docx` en la raíz del repositorio.")
else:
    # Formulario de entrada
    with st.form("form_evidencias"):
        st.markdown("### 1. Programas y Grupos (Fichas)")
        col1, col2 = st.columns(2)
        
        with col1:
            g1 = st.text_input("Grupo 1 (Ej: TOTF 2 / TOTF 2831201):", value="TOTF 2")
            g2 = st.text_input("Grupo 2 (Ej: TMMI 1):", value="TMMI 1")
        with col2:
            g3 = st.text_input("Grupo 3 (Opcional):", value="")

        st.markdown("---")
        st.markdown("### 2. Registro Fotográfico")
        imagenes_cargadas = st.file_uploader(
            "Carga las fotos de evidencia (JPG, PNG):", 
            type=["jpg", "jpeg", "png"], 
            accept_multiple_files=True
        )

        # El botón DEBE estar dentro del 'with st.form'
        submitted = st.form_submit_button("🚀 Generar Informe con Fotografías")

    # Lógica que se ejecuta al enviar el formulario
    if submitted:
        # Construir la cadena de grupos combinados
        grupos_list = [g.strip() for g in [g1, g2, g3] if g.strip()]
        
        if len(grupos_list) > 1:
            texto_grupos = ", ".join(grupos_list[:-1]) + " y " + grupos_list[-1]
        elif len(grupos_list) == 1:
            texto_grupos = grupos_list[0]
        else:
            texto_grupos = "TOTF y TMMI"

        try:
            doc = Document(PLANTILLA_WORD)

            # Mapeo de reemplazo de variables simples
            reemplazos = {
                "{{Fecha actual}}": fecha_hoy_str,
                "{{G1, G2, G3}}": texto_grupos,
                "{{ G1, G2, G3}}": texto_grupos
            }

            # 1. Reemplazo en párrafos normales
            for p in doc.paragraphs:
                for tag, valor in reemplazos.items():
                    if tag in p.text:
                        p.text = p.text.replace(tag, valor)

            # 2. Insertar las imágenes ANTES de la sección "Elaborado por:"
            for p in doc.paragraphs:
                if "{{Suministrar fotos aquí}}" in p.text:
                    p.text = p.text.replace("{{Suministrar fotos aquí}}", "") # Limpiar etiqueta
                    
                    if imagenes_cargadas:
                        num_fotos = len(imagenes_cargadas)
                        filas = (num_fotos + 1) // 2
                        
                        tabla_fotos = doc.add_table(rows=filas, cols=2)
                        tabla_fotos.autofit = False

                        for idx, img_file in enumerate(imagenes_cargadas):
                            fila_idx = idx // 2
                            col_idx = idx % 2
                            
                            celda = tabla_fotos.cell(fila_idx, col_idx)
                            p_celda = celda.paragraphs[0]
                            p_celda.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            
                            img_bytes = io.BytesIO(img_file.read())
                            run = p_celda.add_run()
                            run.add_picture(img_bytes, width=Cm(6.0))
                            
                            p_leyenda = celda.add_paragraph(f"Evidencia {idx + 1}: Formación Práctica")
                            p_leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                        # Colocar la tabla justo antes del párrafo
                        p._p.addprevious(tabla_fotos._tbl)
                    break

            # Guardar archivo generado en memoria
            output = io.BytesIO()
            doc.save(output)
            output.seek(0)

            st.download_button(
                label="📥 Descargar Informe con Evidencias (.docx)",
                data=output,
                file_name=f"Informe_Evidencias_{texto_grupos.replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            st.success("¡Informe generado con éxito con las fotografías ubicadas antes de la firma!")

        except Exception as e:
            st.error(f"Ocurrió un error al procesar el informe: {e}")
